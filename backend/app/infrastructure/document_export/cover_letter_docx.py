"""Structured cover-letter DOCX renderer (TODO T-01).

Refactored from a one-off German Bewerbung script into a data-driven,
locale-aware formal-letter renderer (roughly DIN-5008 order):

sender block → rule → recipient → date → subject → salutation
→ body → closing → signature.

- Input is a flat document dict (same shape the export route already
  passes) plus OPTIONAL structured overrides - no PII defaults, no
  hardcoded employer, dynamic date.
- Returns bytes via BytesIO so the export route can stream it.
- Original implementation; layout inspired by publicly observable
  formal-letter structure only (T-00 decision).
"""

from __future__ import annotations

from datetime import date
from io import BytesIO

from pydantic import BaseModel, Field

MARGIN_TOP_BOTTOM_IN = 0.6
MARGIN_LEFT_RIGHT_IN = 0.8
FONT_NAME = "Calibri"
FONT_SIZE_PT = 10.5
FONT_COLOR = (0x33, 0x33, 0x33)
SENDER_NAME_SIZE_PT = 12
SENDER_NAME_COLOR = (0x11, 0x11, 0x11)
SENDER_CONTACT_SIZE_PT = 9.5
SENDER_CONTACT_COLOR = (0x55, 0x55, 0x55)
SUBJECT_SIZE_PT = 11
SUBJECT_COLOR = (0x00, 0x33, 0x66)
RULE_COLOR = "CCCCCC"


class CoverLetterData(BaseModel):
    sender_name: str = "Candidate"
    sender_street: str = ""
    sender_city: str = ""
    sender_phone: str = ""
    sender_email: str = ""
    recipient_lines: list[str] = Field(default_factory=list)
    company: str = ""
    target_role: str = "role"
    subject: str = ""
    reference_no: str = ""
    salutation: str = ""
    body_paragraphs: list[str] = Field(default_factory=list)
    closing: str = ""
    signature: str = ""
    letter_date: str = ""

    locale: str = "en"


LOCALE_DEFAULTS: dict[str, dict[str, str]] = {
    "de": {"salutation": "Sehr geehrte Damen und Herren,", "closing": "Mit freundlichen Grüßen"},
    "en": {"salutation": "Dear Hiring Manager,", "closing": "Sincerely,"},
    "fr": {"salutation": "Madame, Monsieur,", "closing": "Cordialement,"},
}


def _today_iso() -> str:
    return date.today().isoformat()


def from_document(document: dict) -> CoverLetterData:
    """Build letter data from the flat document dict the API already uses.

    Structured keys win when present; otherwise derive a classic fallback
    from name/company/role/summary so old callers keep working.
    """
    locale = str(document.get("locale") or document.get("language") or "en").lower()
    if locale not in LOCALE_DEFAULTS:
        locale = "en"
    defaults = LOCALE_DEFAULTS[locale]

    sender_name = str(document.get("sender_name") or document.get("name") or "Candidate")
    company = str(document.get("company") or "Target Company")
    role = str(document.get("target_role") or document.get("role") or "role")
    summary = str(document.get("summary") or "")
    body = list(document.get("body_paragraphs") or [])
    if not body:
        if summary:
            body = [summary, "I would welcome the opportunity to contribute."]
        else:
            body = [f"I am excited to apply for the {role} position at {company}."]

    recipient = list(document.get("recipient_lines") or [])
    if not recipient and company and company != "Target Company":
        recipient = [company]

    subject = str(document.get("subject") or f"Application for {role} at {company}")
    reference_no = str(document.get("reference_no") or "")
    if reference_no and reference_no not in subject:
        subject = f"{subject} - Ref: {reference_no}"

    return CoverLetterData(
        sender_name=sender_name,
        sender_street=str(document.get("sender_street") or ""),
        sender_city=str(document.get("sender_city") or ""),
        sender_phone=str(document.get("sender_phone") or "").strip(),
        sender_email=str(document.get("sender_email") or "").strip(),
        recipient_lines=recipient,
        company=company,
        target_role=role,
        subject=subject,
        reference_no=reference_no,
        salutation=str(document.get("salutation") or defaults["salutation"]),
        body_paragraphs=body,
        closing=str(document.get("closing") or defaults["closing"]),
        signature=str(document.get("signature") or sender_name),
        letter_date=str(document.get("letter_date") or _today_iso()),
        locale=locale,
    )


def render_cover_letter_to_docx(document: dict) -> bytes:
    """Render a structured formal cover letter to DOCX bytes."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import parse_xml
    from docx.shared import Inches, Pt, RGBColor

    data = from_document(document)
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(MARGIN_TOP_BOTTOM_IN)
        section.bottom_margin = Inches(MARGIN_TOP_BOTTOM_IN)
        section.left_margin = Inches(MARGIN_LEFT_RIGHT_IN)
        section.right_margin = Inches(MARGIN_LEFT_RIGHT_IN)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(FONT_SIZE_PT)
    normal.font.color.rgb = RGBColor(*FONT_COLOR)

    def spaced(paragraph: object, before: int = 0, after: int = 4) -> None:
        fmt = paragraph.paragraph_format  # type: ignore[attr-defined]
        fmt.space_before = Pt(before)
        fmt.space_after = Pt(after)
        fmt.line_spacing = 1.15

    sender = doc.add_paragraph()
    spaced(sender, before=0, after=2)
    name_run = sender.add_run(data.sender_name + "\n")
    name_run.bold = True
    name_run.font.size = Pt(SENDER_NAME_SIZE_PT)
    name_run.font.color.rgb = RGBColor(*SENDER_NAME_COLOR)
    contact_parts = []
    street_city = " ".join(p for p in [data.sender_street, data.sender_city] if p)
    if street_city:
        contact_parts.append(street_city)
    phone_email = " | ".join(
        p
        for p in [
            f"Phone: {data.sender_phone}" if data.sender_phone else "",
            f"Email: {data.sender_email}" if data.sender_email else "",
        ]
        if p
    )
    if phone_email:
        contact_parts.append(phone_email)
    if contact_parts:
        contact_run = sender.add_run("\n".join(contact_parts))
        contact_run.font.size = Pt(SENDER_CONTACT_SIZE_PT)
        contact_run.font.color.rgb = RGBColor(*SENDER_CONTACT_COLOR)

    rule = doc.add_paragraph()
    spaced(rule, before=2, after=12)
    border = parse_xml(
        '<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:bottom w:val="single" w:sz="6" w:space="1" w:color="{RULE_COLOR}"/>'
        "</w:pBdr>"
    )
    rule._p.get_or_add_pPr().append(border)

    if data.recipient_lines:
        recipient = doc.add_paragraph()
        spaced(recipient, before=0, after=12)
        recipient.add_run("\n".join(data.recipient_lines))

    date_par = doc.add_paragraph()
    date_par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    spaced(date_par, before=0, after=14)
    date_par.add_run(data.letter_date)

    subject = doc.add_paragraph()
    spaced(subject, before=0, after=12)
    subject_run = subject.add_run(data.subject)
    subject_run.bold = True
    subject_run.font.size = Pt(SUBJECT_SIZE_PT)
    subject_run.font.color.rgb = RGBColor(*SUBJECT_COLOR)

    salutation = doc.add_paragraph()
    spaced(salutation, before=0, after=8)
    salutation.add_run(data.salutation)

    for text in data.body_paragraphs:
        body = doc.add_paragraph()
        spaced(body, before=0, after=8)
        body.add_run(text)

    closing = doc.add_paragraph()
    spaced(closing, before=0, after=18)
    closing.add_run(data.closing)
    signature = doc.add_paragraph()
    spaced(signature, before=0, after=0)
    signature_run = signature.add_run(data.signature or data.sender_name)
    signature_run.bold = True

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
