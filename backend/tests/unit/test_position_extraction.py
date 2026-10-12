"""Position extraction from text, HTML, PDF, DOCX, images, ZIP and URLs."""

import io
import json
import shutil
import zipfile

import httpx
import pytest
from docx import Document
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from app.application.services import position_extraction as px
from app.core.errors import ConfigurationError, ExternalServiceError, ValidationError

AD_EN = """Senior Python Developer (m/w/d)
Company: Acme GmbH
Location: Berlin

We are looking for an experienced Python developer to build FastAPI services.
Requirements: Python, FastAPI, PostgreSQL, Docker and strong communication.
You will design APIs, review code and mentor the team. Python experience is essential.
"""

AD_DE = """Stellenbezeichnung: Softwareentwickler Backend
Unternehmen: Müller & Söhne AG
Standort: Köln

Wir suchen eine erfahrene Entwicklerin für unsere Backend-Plattform. Sie arbeiten mit Python,
Kubernetes und PostgreSQL. Gute Deutschkenntnisse (Umlaute: ä ö ü ß) sind erforderlich.
"""


def _pdf(text: str) -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    y = 800
    for line in text.splitlines():
        pdf.drawString(50, y, line)
        y -= 16
    pdf.save()
    return buffer.getvalue()


def _docx(text: str) -> bytes:
    document = Document()
    for line in text.splitlines():
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _png(lines: list[str]) -> bytes:
    image = Image.new("RGB", (1000, 60 + 60 * len(lines)), "white")
    draw = ImageDraw.Draw(image)
    # Render with a real vector font so Tesseract gets legible glyphs. The
    # original DejaVu path only exists on Linux; macOS/Windows runners fall back
    # to a tiny bitmap font that OCR misreads, so try common per-OS locations
    # before giving up. This keeps the "real tesseract" test meaningful
    # everywhere instead of silently degrading image quality.
    font = None
    for candidate in (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Helvetica.ttf",
        "/Library/Fonts/Arial Unicode.ttf",
        "C:/Windows/Fonts/arial.ttf",
    ):
        try:
            font = ImageFont.truetype(candidate, 32)
            break
        except OSError:
            continue
    if font is None:  # pragma: no cover - only on exotic runners
        font = ImageFont.load_default()
    for index, line in enumerate(lines):
        draw.text((20, 20 + 60 * index), line, fill="black", font=font)
    buffer = io.BytesIO()
    image.save(buffer, "PNG")
    return buffer.getvalue()


def _zip(files: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return buffer.getvalue()


class TestTextParsing:
    def test_labelled_english_ad(self):
        position = px.extract_from_text(AD_EN).positions[0]
        assert position.title == "Senior Python Developer (m/w/d)"
        assert position.company == "Acme GmbH"
        assert position.location == "Berlin"
        assert position.source == "text"
        assert "python" in position.keywords
        assert position.warnings == ()

    def test_labelled_german_ad_keeps_umlauts(self):
        position = px.extract_from_text(AD_DE).positions[0]
        assert position.title == "Softwareentwickler Backend"
        assert position.company == "Müller & Söhne AG"
        assert position.location == "Köln"
        assert "ä ö ü ß" in position.description

    def test_title_at_company_pattern(self):
        text = "Data Engineer at Globex Corp\n\n" + "We process large volumes of data every day. " * 3
        position = px.extract_from_text(text).positions[0]
        assert (position.title, position.company) == ("Data Engineer", "Globex Corp")
        assert "location could not be determined" in position.warnings

    def test_missing_fields_are_none_with_warnings_never_invented(self):
        text = "Join our team and help us build great products for customers all around the world today."
        position = px.extract_from_text(text).positions[0]
        assert position.company is None
        assert position.location is None
        assert {"company could not be determined", "location could not be determined"} <= set(position.warnings)

    @pytest.mark.parametrize("text", ["", "   ", "short"])
    def test_rejects_empty_or_too_short_text(self, text):
        with pytest.raises(ValidationError):
            px.extract_from_text(text)

    def test_text_is_capped(self):
        position = px.extract_from_text("Title: Dev\n" + "word " * 40_000).positions[0]
        assert len(position.description) <= px.MAX_TEXT_CHARS

    def test_keywords_drop_stopwords_and_are_deterministic(self):
        keywords = px.extract_keywords("the python and the python und die FastAPI FastAPI FastAPI docker")
        assert keywords[0] == "fastapi"
        assert "the" not in keywords and "und" not in keywords
        assert keywords == px.extract_keywords("the python and the python und die FastAPI FastAPI FastAPI docker")


class TestHtml:
    def test_json_ld_job_posting_is_preferred(self):
        posting = {
            "@context": "https://schema.org",
            "@type": "JobPosting",
            "title": "Platform Engineer",
            "hiringOrganization": {"@type": "Organization", "name": "Initech"},
            "jobLocation": {"@type": "Place", "address": {"addressLocality": "Hamburg", "addressCountry": "DE"}},
            "datePosted": "2026-10-01",
            "employmentType": ["FULL_TIME"],
            "description": "<p>Build <b>Kubernetes</b> platforms and CI/CD pipelines for all engineering teams.</p>",
        }
        html = (
            "<html><head><title>x</title>"
            f"<script type='application/ld+json'>{json.dumps(posting)}</script>"
            "</head><body>noise</body></html>"
        )
        position = px.extract_from_text(html).positions[0]
        assert position.extracted_with == "json_ld"
        assert (position.title, position.company, position.location) == ("Platform Engineer", "Initech", "Hamburg, DE")
        assert position.published_at == "2026-10-01"
        assert position.employment_type == "FULL_TIME"
        assert "Kubernetes" in position.description and "<b>" not in position.description

    def test_json_ld_inside_graph(self):
        graph = {
            "@graph": [
                {"@type": "WebSite"},
                {"@type": "JobPosting", "title": "QA Lead", "description": "Lead testing " * 5},
            ]
        }
        html = f"<html><script type='application/ld+json'>{json.dumps(graph)}</script><body>x</body></html>"
        assert px.extract_from_text(html).positions[0].title == "QA Lead"

    def test_invalid_json_ld_falls_back_to_visible_text(self):
        html = (
            "<html><head><title>Backend Developer</title><script type='application/ld+json'>{broken</script></head>"
            "<body><h1>Backend Developer</h1><p>Company: Hooli</p>"
            "<p>We need a developer who loves distributed systems.</p></body></html>"
        )
        position = px.extract_from_text(html).positions[0]
        assert position.extracted_with == "heuristic"
        assert position.company == "Hooli"

    def test_scripts_nav_and_footer_are_not_part_of_the_description(self):
        html = (
            "<html><body><nav>MENU-LINKS</nav><script>var secret=1</script><h1>Dev</h1>"
            "<p>Real job content about building reliable software for customers.</p>"
            "<footer>COOKIES</footer></body></html>"
        )
        description = px.extract_from_text(html).positions[0].description
        assert "MENU-LINKS" not in description and "secret" not in description and "COOKIES" not in description
        assert "Real job content" in description


class TestFiles:
    def test_pdf(self):
        position = px.extract_from_file(_pdf(AD_EN), filename="ad.pdf").positions[0]
        assert position.source == "pdf"
        assert position.company == "Acme GmbH"

    def test_docx(self):
        position = px.extract_from_file(_docx(AD_DE), filename="ad.docx").positions[0]
        assert position.source == "docx"
        assert position.title == "Softwareentwickler Backend"

    def test_docx_tables_are_read(self):
        document = Document()
        document.add_paragraph("Job Title: Analyst")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Requirements"
        table.rows[0].cells[1].text = "SQL, Excel and clear communication skills for stakeholders"
        buffer = io.BytesIO()
        document.save(buffer)
        position = px.extract_from_file(buffer.getvalue()).positions[0]
        assert "SQL, Excel" in position.description

    def test_plain_text_and_html_files(self):
        assert px.extract_from_file(AD_EN.encode()).positions[0].source == "text"
        assert px.extract_from_file(f"<html><body><p>{AD_EN}</p></body></html>".encode()).positions[0].source == "html"

    def test_latin1_text_is_decoded(self):
        position = px.extract_from_file(AD_DE.encode("cp1252", errors="replace")).positions[0]
        assert position.company == "Müller & Söhne AG"
        assert "ä ö ü ß" in position.description

    def test_utf16_is_decoded_only_with_a_bom(self):
        position = px.extract_from_file(AD_EN.encode("utf-16")).positions[0]
        assert position.company == "Acme GmbH"

    def test_type_is_detected_from_bytes_not_from_the_name(self):
        assert px.extract_from_file(_pdf(AD_EN), filename="harmless.txt").positions[0].source == "pdf"

    def test_corrupt_pdf_is_a_clean_validation_error(self):
        with pytest.raises(ValidationError):
            px.extract_from_file(b"%PDF-1.7\nthis is not a pdf")

    def test_corrupt_docx_is_a_clean_validation_error(self):
        with pytest.raises(ValidationError):
            px.extract_from_file(b"PK\x03\x04garbage-garbage-garbage")

    def test_binary_blob_is_rejected(self):
        with pytest.raises(ValidationError, match="Unsupported file type"):
            px.extract_from_file(b"\x00\x01\x02\x03" * 100)

    def test_empty_and_oversized_uploads_are_rejected(self):
        with pytest.raises(ValidationError, match="empty"):
            px.extract_from_file(b"")
        with pytest.raises(ValidationError, match="larger"):
            px.extract_from_file(b"a" * (px.MAX_UPLOAD_BYTES + 1))


class TestImages:
    def test_image_uses_the_injected_ocr_engine(self):
        seen = {}

        def fake_ocr(data: bytes) -> str:
            seen["bytes"] = len(data)
            return AD_EN

        position = px.extract_from_file(_png(["x"]), ocr=fake_ocr).positions[0]
        assert position.source == "image"
        assert position.company == "Acme GmbH"
        assert seen["bytes"] > 0

    def test_missing_ocr_engine_is_a_configuration_error(self, monkeypatch):
        monkeypatch.setattr(px.shutil, "which", lambda name: None)
        with pytest.raises(ConfigurationError, match="OCR"):
            px.extract_from_file(_png(["Senior Python Developer"]))

    def test_corrupt_image_is_rejected_before_ocr(self):
        called = []
        with pytest.raises(ValidationError):
            px.extract_from_file(b"\x89PNG\r\n\x1a\nnot really a png", ocr=lambda data: called.append(1) or "x")
        assert not called

    def test_huge_resolution_is_rejected_before_ocr(self, monkeypatch):
        monkeypatch.setattr(px, "MAX_IMAGE_PIXELS", 100)
        with pytest.raises(ValidationError, match="resolution"):
            px.extract_from_file(_png(["Hello"]), ocr=lambda data: AD_EN)

    @pytest.mark.skipif(shutil.which("tesseract") is None, reason="tesseract is not installed")
    def test_real_tesseract_reads_a_screenshot(self):
        image = _png(
            [
                "Senior Python Developer (m/w/d)",
                "Company: Acme GmbH",
                "Location: Berlin",
                "Requirements: Python FastAPI Docker PostgreSQL",
            ]
        )
        position = px.extract_from_file(image).positions[0]
        assert "Python" in position.description
        assert position.company and "Acme" in position.company


class TestZip:
    def test_extracts_each_supported_member_and_reports_the_rest(self):
        archive = _zip(
            {
                "ads/backend.pdf": _pdf(AD_EN),
                "ads/frontend.docx": _docx(AD_DE),
                "notes.txt": AD_EN.encode(),
                "photo.exe": b"MZ" + b"\x00" * 10,
                "__MACOSX/._backend.pdf": b"junk",
                ".hidden.txt": b"hidden",
            }
        )
        result = px.extract_from_file(archive, filename="ads.zip")
        assert [p.source for p in result.positions] == ["zip:backend.pdf", "zip:frontend.docx", "zip:notes.txt"]
        assert {item["name"]: item["reason"] for item in result.skipped} == {"photo.exe": "unsupported file type"}

    def test_path_traversal_names_are_skipped_not_followed(self):
        result = px.extract_from_file(_zip({"../../etc/passwd.txt": AD_EN.encode(), "ok.txt": AD_EN.encode()}))
        assert [p.source for p in result.positions] == ["zip:ok.txt"]
        assert result.skipped == [{"name": "../../etc/passwd.txt", "reason": "unsafe path"}]

    def test_nested_archives_are_not_opened(self):
        inner = _zip({"a.txt": AD_EN.encode()})
        result = px.extract_from_file(_zip({"inner.zip": inner, "ok.txt": AD_EN.encode()}))
        assert [p.source for p in result.positions] == ["zip:ok.txt"]
        # inner.zip is unsupported by extension before it is ever opened
        assert result.skipped[0]["name"] == "inner.zip"

    def test_zip_bomb_is_rejected(self):
        bomb = _zip({"big.txt": b"A" * (30 * 1024 * 1024)})
        assert len(bomb) < px.MAX_UPLOAD_BYTES  # small upload, huge expansion
        with pytest.raises(ValidationError):
            px.extract_from_file(bomb)

    def test_too_many_entries_are_rejected(self):
        many = _zip({f"f{i}.txt": AD_EN.encode() for i in range(px.MAX_ZIP_ENTRIES + 1)})
        with pytest.raises(ValidationError, match="more than"):
            px.extract_from_file(many)

    def test_one_bad_member_does_not_fail_the_whole_archive(self):
        result = px.extract_from_file(_zip({"bad.pdf": b"%PDF-1.4 broken", "good.txt": AD_EN.encode()}))
        assert [p.source for p in result.positions] == ["zip:good.txt"]
        assert result.skipped[0]["name"] == "bad.pdf"

    def test_archive_without_files_is_rejected(self):
        with pytest.raises(ValidationError):
            px.extract_from_file(_zip({}))


PUBLIC = lambda host: ["93.184.216.34"]  # noqa: E731 - a public documentation address


def _transport(handler):
    return httpx.MockTransport(handler)


class TestUrlFetching:
    def test_fetches_and_extracts_with_pinned_ip_and_original_host(self):
        seen = {}

        def handler(request: httpx.Request) -> httpx.Response:
            seen["url"] = str(request.url)
            seen["host"] = request.headers["host"]
            return httpx.Response(
                200, headers={"content-type": "text/html"}, text=f"<html><body><p>{AD_EN}</p></body></html>"
            )

        result = px.extract_from_url(
            "https://jobs.example/posting/42?utm_source=x", resolver=PUBLIC, transport=_transport(handler)
        )
        position = result.positions[0]
        assert position.source == "url"
        assert position.url == "https://jobs.example/posting/42?utm_source=x"
        assert position.company == "Acme GmbH"
        assert seen["url"].startswith("https://93.184.216.34/posting/42")  # connection pinned to the checked IP
        assert seen["host"] == "jobs.example"  # TLS/HTTP still address the real host

    @pytest.mark.parametrize(
        "url",
        [
            "file:///etc/passwd",
            "ftp://jobs.example/a",
            "http://user:pass@jobs.example/a",
            "https://jobs.example:8443/a",
            "https:///nohost",
            "http://127.0.0.1/admin",
            "http://169.254.169.254/latest/meta-data/",
            "http://10.0.0.5/",
            "http://192.168.1.1/",
            "http://100.64.0.1/",
            "http://[::1]/",
            "http://[::ffff:127.0.0.1]/",
            "http://0.0.0.0/",
        ],
    )
    def test_ssrf_targets_are_rejected_without_any_request(self, url):
        def handler(request):  # pragma: no cover - must never run
            raise AssertionError("no HTTP request may be made for a blocked URL")

        with pytest.raises(ValidationError):
            px.extract_from_url(url, resolver=PUBLIC, transport=_transport(handler))

    def test_hostname_resolving_to_a_private_address_is_rejected(self):
        with pytest.raises(ValidationError, match="non-public"):
            px.extract_from_url(
                "https://evil.example/a",
                resolver=lambda host: ["10.1.2.3"],
                transport=_transport(lambda r: httpx.Response(200)),
            )

    def test_any_private_address_among_several_blocks_the_host(self):
        with pytest.raises(ValidationError, match="non-public"):
            px.extract_from_url(
                "https://mixed.example/",
                resolver=lambda host: ["93.184.216.34", "127.0.0.1"],
                transport=_transport(lambda r: httpx.Response(200)),
            )

    def test_unresolvable_host_is_a_validation_error(self):
        with pytest.raises(ValidationError, match="resolved"):
            px.extract_from_url(
                "https://nope.example/", resolver=lambda host: [], transport=_transport(lambda r: httpx.Response(200))
            )

    def test_redirect_to_a_private_address_is_blocked(self):
        def handler(request):
            return httpx.Response(302, headers={"location": "http://169.254.169.254/latest/meta-data/"})

        with pytest.raises(ValidationError, match="non-public"):
            px.extract_from_url(
                "https://jobs.example/r",
                resolver=lambda host: ["93.184.216.34"] if "jobs" in host else ["169.254.169.254"],
                transport=_transport(handler),
            )

    def test_public_redirect_is_followed_and_revalidated(self):
        hosts = []

        def handler(request):
            hosts.append(request.headers["host"])
            if request.headers["host"] == "short.example":
                return httpx.Response(301, headers={"location": "https://jobs.example/real"})
            return httpx.Response(200, headers={"content-type": "text/plain"}, text=AD_EN)

        result = px.extract_from_url("https://short.example/x", resolver=PUBLIC, transport=_transport(handler))
        assert hosts == ["short.example", "jobs.example"]
        assert result.positions[0].url == "https://jobs.example/real"

    def test_redirect_loops_are_bounded(self):
        with pytest.raises(ValidationError, match="too many times"):
            px.extract_from_url(
                "https://a.example/",
                resolver=PUBLIC,
                transport=_transport(lambda r: httpx.Response(302, headers={"location": "https://a.example/"})),
            )

    def test_http_error_status_is_an_external_service_error(self):
        with pytest.raises(ExternalServiceError, match="HTTP 404"):
            px.extract_from_url(
                "https://jobs.example/gone", resolver=PUBLIC, transport=_transport(lambda r: httpx.Response(404))
            )

    def test_oversized_response_is_rejected(self):
        body = b"x" * (px.MAX_FETCH_BYTES + 10)
        with pytest.raises(ValidationError, match="too large"):
            px.extract_from_url(
                "https://jobs.example/big",
                resolver=PUBLIC,
                transport=_transport(lambda r: httpx.Response(200, content=body)),
            )

    def test_declared_oversize_is_rejected_early(self):
        response = httpx.Response(200, headers={"content-length": str(px.MAX_FETCH_BYTES + 1)}, content=b"x")
        with pytest.raises(ValidationError, match="too large"):
            px.extract_from_url("https://jobs.example/big", resolver=PUBLIC, transport=_transport(lambda r: response))

    def test_network_errors_do_not_leak_details(self):
        def handler(request):
            raise httpx.ConnectError("secret internal detail 10.0.0.9")

        with pytest.raises(ExternalServiceError) as caught:
            px.extract_from_url("https://jobs.example/x", resolver=PUBLIC, transport=_transport(handler))
        assert "10.0.0.9" not in caught.value.message

    def test_pdf_served_from_a_url_is_parsed(self):
        result = px.extract_from_url(
            "https://jobs.example/ad.pdf",
            resolver=PUBLIC,
            transport=_transport(lambda r: httpx.Response(200, content=_pdf(AD_EN))),
        )
        assert result.positions[0].company == "Acme GmbH"

    def test_archives_behind_a_url_are_refused(self):
        with pytest.raises(ValidationError, match="job posting page"):
            px.extract_from_url(
                "https://jobs.example/a.zip",
                resolver=PUBLIC,
                transport=_transport(lambda r: httpx.Response(200, content=_zip({"a.txt": b"x"}))),
            )
