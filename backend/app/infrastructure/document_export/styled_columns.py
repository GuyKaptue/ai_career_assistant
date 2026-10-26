"""Two-column CV layouts: sidebar left or right, adjustable width, optional full-width name band.

The sidebar holds the compact reference sections (skills, languages, certifications); the main
column holds profile, projects, experience and education. Colours come from the selected
``CvTheme``; ``theme.sidebar_side`` / ``sidebar_pct`` / ``top_band`` / ``heading`` describe the
design. An initials avatar is drawn only when ``theme.avatar`` is set (samples); real photos are
never stored or generated here.
"""

from __future__ import annotations

import textwrap
from html import escape
from io import BytesIO

from app.infrastructure.document_export.styled import Parsed, Section, Theme, _split_bullet, parse_content

_SIDEBAR_KEYS = {
    "skills",
    "kenntnisse",
    "compétences",
    "languages",
    "sprachen",
    "langues",
    "certifications",
    "zertifikate",
    "certificats",
    "tools",
}
SIDEBAR_PCT = 32
BAND_H = 92  # PDF points


def split_sections(p: Parsed) -> tuple[list[Section], list[Section]]:
    side = [s for s in p.sections if s.heading and s.heading.lower() in _SIDEBAR_KEYS]
    main = [s for s in p.sections if s not in side]
    return side, main


def on_color(background: str) -> str:
    """Black or white, whichever reads better on ``background``."""
    from app.domain.templates.themes import contrast_ratio

    return "#FFFFFF" if contrast_ratio("#FFFFFF", background) >= contrast_ratio("#000000", background) else "#000000"


def initials(name: str) -> str:
    parts = [w for w in name.split() if w[:1].isalpha()]
    return "".join(w[0].upper() for w in parts[:2]) or "?"


def _font(theme: Theme) -> tuple[str, str]:
    return ("Times-Roman", "Times-Bold") if theme.font == "Times-Roman" else ("Helvetica", "Helvetica-Bold")


def _pal(theme: Theme):
    if theme.palette is None:
        raise ValueError("theme.palette is required for two-column rendering")
    return theme.palette


def _esc(t: str) -> str:
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


# ----------------------------------------------------------------------------- PDF


def pdf_two_column(title: str, content: str, theme: Theme) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import (
        BaseDocTemplate,
        Frame,
        FrameBreak,
        HRFlowable,
        NextPageTemplate,
        PageTemplate,
        Paragraph,
        Spacer,
    )

    pal, (base, bold) = _pal(theme), _font(theme)
    p = parse_content(content, "cv", title)
    side, main = split_sections(p)
    w, h = A4
    sw = w * theme.sidebar_pct / 100
    sx = 0 if theme.sidebar_side == "left" else w - sw
    mx = sw if theme.sidebar_side == "left" else 0
    pad = 16
    hx = colors.HexColor
    band = BAND_H if theme.top_band else 0
    band_fg = on_color(pal.accent)

    def background(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(hx(pal.background))
        canvas.rect(0, 0, w, h, stroke=0, fill=1)
        if doc.page == 1:
            canvas.setFillColor(hx(pal.sidebar))
            canvas.rect(sx, 0, sw, h - band, stroke=0, fill=1)
            if band:
                canvas.setFillColor(hx(pal.accent))
                canvas.rect(0, h - band, w, band, stroke=0, fill=1)
                canvas.setFillColor(hx(band_fg))
                canvas.setFont(bold, 26)
                canvas.drawString(34, h - 46, p.name)
                if p.subtitle:
                    canvas.setFont(base, 12)
                    canvas.drawString(34, h - 66, p.subtitle)
            if theme.avatar:
                cx, cy, r = (w - 62, h - band / 2, 30) if band else (sx + sw / 2, h - 70, 35)
                canvas.setStrokeColor(hx(band_fg if band else pal.accent))
                canvas.setLineWidth(2)
                canvas.setFillColor(hx(pal.sidebar if band else pal.accent))
                canvas.circle(cx, cy, r, stroke=1, fill=1)
                canvas.setFillColor(hx(on_color(pal.sidebar if band else pal.accent)))
                canvas.setFont(bold, 20)
                canvas.drawCentredString(cx, cy - 7, initials(p.name))
        canvas.restoreState()

    side_top = band + 20 if band else (120 if theme.avatar else 40)
    first = [
        Frame(sx, 30, sw, h - side_top - 30, leftPadding=pad, rightPadding=pad, topPadding=0, id="side"),
        Frame(
            mx,
            30,
            w - sw,
            h - band - 50 if band else h - 60,
            leftPadding=pad + 4,
            rightPadding=pad + 4,
            topPadding=0,
            id="main",
        ),
    ]
    later = [Frame(30, 30, w - 60, h - 60, id="later")]
    buf = BytesIO()
    doc = BaseDocTemplate(buf, pagesize=A4, title=title or p.name)
    doc.addPageTemplates(
        [
            PageTemplate(id="first", frames=first, onPage=background),
            PageTemplate(id="later", frames=later, onPage=background),
        ]
    )
    ss = ParagraphStyle(
        "sh", fontName=bold, fontSize=10.5, leading=13, textColor=hx(pal.sidebar_heading), spaceBefore=12, spaceAfter=4
    )
    sb = ParagraphStyle("sb", fontName=base, fontSize=9, leading=12, textColor=hx(pal.sidebar_text), spaceAfter=2)
    name_s = ParagraphStyle("n", fontName=bold, fontSize=24, leading=28, textColor=hx(pal.section_title))
    role_s = ParagraphStyle("r", fontName=base, fontSize=12, leading=15, textColor=hx(pal.section_title), spaceAfter=6)
    ms = ParagraphStyle(
        "mh", fontName=bold, fontSize=11.5, leading=14, textColor=hx(pal.section_title), spaceBefore=12, spaceAfter=3
    )
    mb = ParagraphStyle("mb", fontName=base, fontSize=10, leading=13.5, textColor=hx(pal.text), spaceAfter=3)
    m1 = ParagraphStyle("m1", parent=mb, leftIndent=12, bulletIndent=2, spaceAfter=2)
    m2 = ParagraphStyle("m2", parent=m1, leftIndent=26, bulletIndent=16)

    story: list = [NextPageTemplate("later")]
    for sec in side:
        story.append(Paragraph(_esc((sec.heading or "").upper()), ss))
        for ln in sec.lines:
            b = _split_bullet(ln)
            text = b[1] if b else ln.strip()
            if text:
                story.append(Paragraph(_esc(text), sb, bulletText="•" if b else None))
    story.append(FrameBreak())
    if not band:
        story.append(Paragraph(_esc(p.name), name_s))
        if p.subtitle:
            story.append(Paragraph(_esc(p.subtitle), role_s))
    for sec in main:
        if sec.heading:
            head = sec.heading.upper() if theme.heading in ("caps", "rule", "bar") else sec.heading
            story.append(Paragraph(_esc(head), ms))
            if theme.heading == "rule":
                story.append(HRFlowable(width="100%", thickness=1.2, color=hx(pal.accent), spaceAfter=3))
            elif theme.heading == "bar":
                story.append(HRFlowable(width=34, thickness=3, color=hx(pal.accent), spaceAfter=3, hAlign="LEFT"))
        for ln in sec.lines:
            b = _split_bullet(ln)
            if b:
                story.append(Paragraph(_esc(b[1]), m2 if b[0] else m1, bulletText="•"))
            elif ln.strip():
                story.append(Paragraph(_esc(ln.strip()), mb))
            else:
                story.append(Spacer(1, 3))
    doc.build(story)
    return buf.getvalue()


# ----------------------------------------------------------------------------- HTML


def _html_sections(sections: list[Section], heading_class: str) -> str:
    out: list[str] = []
    for sec in sections:
        if sec.heading:
            out.append(f'<h2 class="{heading_class}">{escape(sec.heading)}</h2>')
        items: list[str] = []
        for ln in sec.lines:
            b = _split_bullet(ln)
            if b:
                items.append(f'<li class="l{b[0]}">{escape(b[1])}</li>')
                continue
            if items:
                out.append("<ul>" + "".join(items) + "</ul>")
                items = []
            if ln.strip():
                out.append(f"<p>{escape(ln.strip())}</p>")
        if items:
            out.append("<ul>" + "".join(items) + "</ul>")
    return "".join(out)


def html_two_column(title: str, content: str, theme: Theme) -> str:
    pal = _pal(theme)
    p = parse_content(content, "cv", title)
    side, main = split_sections(p)
    right = theme.sidebar_side == "right"
    family = (
        "Georgia,'Times New Roman',serif" if theme.font == "Times-Roman" else "Inter,'Helvetica Neue',Arial,sans-serif"
    )
    avatar = f'<div class="av">{escape(initials(p.name))}</div>' if theme.avatar else ""
    sub = f'<div class="role">{escape(p.subtitle)}</div>' if p.subtitle else ""
    pct = theme.sidebar_pct
    cols = f"1fr {pct}%" if right else f"{pct}% 1fr"
    band_fg = on_color(pal.accent)
    rule = {
        "rule": f"border-bottom:2px solid {pal.accent};padding-bottom:3px;",
        "bar": f"border-left:5px solid {pal.accent};padding-left:8px;",
    }.get(theme.heading, "")
    band_html = f'<header class="band"><div><h1>{escape(p.name)}</h1>{sub}</div></header>' if theme.top_band else ""
    name_html = "" if theme.top_band else f"<h1>{escape(p.name)}</h1>{sub}"
    aside_avatar = "" if theme.top_band else avatar
    css = (
        f"body{{margin:0;background:#e5e7eb;font-family:{family};font-size:14px;line-height:1.45}}"
        f".page{{display:grid;grid-template-columns:{cols};max-width:794px;min-height:1123px;"
        f"margin:24px auto;background:{pal.background};color:{pal.text};box-shadow:0 1px 6px #0002}}"
        f".band{{grid-column:1/-1;background:{pal.accent};color:{band_fg};padding:22px 28px}}"
        ".band h1{margin:0;font-size:30px}.band .role{color:inherit;opacity:.9}"
        f".side{{background:{pal.sidebar};color:{pal.sidebar_text};padding:28px 18px;"
        f"grid-column:{2 if right else 1};grid-row:{2 if theme.top_band else 1}}}"
        f".side h2{{color:{pal.sidebar_heading};font-size:13px;letter-spacing:.08em;"
        "text-transform:uppercase;margin:18px 0 6px}"
        f".main{{padding:28px 26px;grid-column:{1 if right else 2};grid-row:{2 if theme.top_band else 1}}}"
        f".main h1{{margin:0;font-size:30px;color:{pal.section_title}}}"
        f".role{{color:{pal.section_title};margin:2px 0 10px}}"
        f".main h2{{color:{pal.section_title};font-size:15px;text-transform:uppercase;letter-spacing:.06em;"
        f"{rule}margin:18px 0 6px}}"
        "ul{margin:4px 0;padding-left:18px}li.l1{margin-left:16px}p{margin:6px 0}"
        f".av{{width:96px;height:96px;border-radius:50%;background:{pal.accent};color:{on_color(pal.accent)};"
        "display:flex;align-items:center;justify-content:center;font-size:30px;font-weight:700;margin:0 auto 12px;"
        f"border:2px solid {pal.accent}}}"
        "@media(max-width:640px){.page{grid-template-columns:1fr}.side,.main{grid-column:1;grid-row:auto}}"
    )
    return (
        '<!DOCTYPE html>\n<html lang="en"><head><meta charset="utf-8" />'
        '<meta name="viewport" content="width=device-width,initial-scale=1" />'
        f"<title>{escape(title or p.name)}</title><style>{css}</style></head><body>"
        f'<div class="page">{band_html}<aside class="side">{aside_avatar}{_html_sections(side, "sh")}</aside>'
        f'<main class="main">{name_html}{_html_sections(main, "mh")}</main></div></body></html>\n'
    )


# ----------------------------------------------------------------------------- DOCX


def docx_two_column(title: str, content: str, theme: Theme) -> bytes:
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor

    from app.infrastructure.document_export.styled import _bottom_border, _shade

    pal = _pal(theme)
    p = parse_content(content, "cv", title)
    side, main = split_sections(p)
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(0.8)
        s.top_margin = s.bottom_margin = Cm(0.8)
    font_name = "Times New Roman" if theme.font == "Times-Roman" else "Calibri"
    doc.styles["Normal"].font.name = font_name
    doc.styles["Normal"].font.size = Pt(10)

    def rgb(hex_color: str):
        return RGBColor.from_string(hex_color.lstrip("#").upper())

    if theme.top_band:
        bt = doc.add_table(rows=1, cols=1)
        bc = bt.rows[0].cells[0]
        _shade(bc, pal.accent)
        run = bc.paragraphs[0].add_run(p.name)
        run.bold, run.font.size, run.font.color.rgb = True, Pt(24), rgb(on_color(pal.accent))
        if p.subtitle:
            sr = bc.add_paragraph().add_run(p.subtitle)
            sr.font.size, sr.font.color.rgb = Pt(11.5), rgb(on_color(pal.accent))

    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    cells = table.rows[0].cells
    right = theme.sidebar_side == "right"
    left_cell, right_cell = (cells[1], cells[0]) if right else (cells[0], cells[1])  # sidebar, main
    total = 19.4
    sidebar_w, main_w = Cm(total * theme.sidebar_pct / 100), Cm(total * (100 - theme.sidebar_pct) / 100)
    left_cell.width, right_cell.width = sidebar_w, main_w
    _shade(left_cell, pal.sidebar)
    _shade(right_cell, pal.background)

    def fill(cell, sections: list[Section], head_color: str, text_color: str, underline: bool) -> None:
        first = True
        for sec in sections:
            if sec.heading:
                para = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                para.paragraph_format.space_before = Pt(8)
                run = para.add_run(sec.heading.upper())
                run.bold, run.font.size, run.font.color.rgb = True, Pt(10.5), rgb(head_color)
                if underline:
                    _bottom_border(para, pal.accent)
            for ln in sec.lines:
                b = _split_bullet(ln)
                text = b[1] if b else ln.strip()
                if not text:
                    continue
                para = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                run = para.add_run(("• " if b else "") + text)
                run.font.size, run.font.color.rgb = Pt(9.5), rgb(text_color)
                if b and b[0]:
                    para.paragraph_format.left_indent = Cm(0.8)

    if theme.avatar:
        ar = left_cell.paragraphs[0].add_run(f"[ {initials(p.name)} ]")
        ar.bold, ar.font.size, ar.font.color.rgb = True, Pt(20), rgb(pal.accent)
        left_cell.add_paragraph()
    fill(left_cell, side, pal.sidebar_heading, pal.sidebar_text, False)
    if not theme.top_band:
        nr = right_cell.paragraphs[0].add_run(p.name)
        nr.bold, nr.font.size, nr.font.color.rgb = True, Pt(22), rgb(pal.section_title)
        if p.subtitle:
            sr = right_cell.add_paragraph().add_run(p.subtitle)
            sr.font.size, sr.font.color.rgb = Pt(11.5), rgb(pal.section_title)
        right_cell.add_paragraph()
    fill(right_cell, main, pal.section_title, pal.text, theme.heading == "rule")
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ----------------------------------------------------------------------------- PNG


def png_two_column(content: str, theme: Theme, width: int = 600) -> bytes:
    from PIL import Image, ImageDraw

    from app.infrastructure.document_export.previews import _pil_font

    pal = _pal(theme)
    height = int(width * 1.414)
    img = Image.new("RGB", (width, height), pal.background)
    d = ImageDraw.Draw(img)
    sw = int(width * theme.sidebar_pct / 100)
    sx0 = 0 if theme.sidebar_side == "left" else width - sw
    band_h = int(width * 0.155) if theme.top_band else 0
    d.rectangle([sx0, band_h, sx0 + sw, height], fill=pal.sidebar)
    p = parse_content(content, "cv")
    side, main = split_sections(p)
    pad = int(width * 0.04)
    name_font, role_font = _pil_font(int(width * 0.05), True), _pil_font(int(width * 0.028))
    if band_h:
        d.rectangle([0, 0, width, band_h], fill=pal.accent)
        fg = on_color(pal.accent)
        d.text((pad, int(width * 0.04)), p.name, font=name_font, fill=fg)
        if p.subtitle:
            d.text((pad, int(width * 0.1)), p.subtitle, font=role_font, fill=fg)
    if theme.avatar:
        r = int(width * 0.075)
        cx = sx0 + sw // 2
        cy = band_h + int(width * 0.115)
        d.ellipse([cx - r, cy - r, cx + r, cy + r], fill=pal.accent)
        f = _pil_font(int(width * 0.04), True)
        t = initials(p.name)
        d.text((cx - d.textlength(t, font=f) / 2, cy - int(width * 0.022)), t, font=f, fill=on_color(pal.accent))
    small, head = _pil_font(int(width * 0.024)), _pil_font(int(width * 0.028), True)
    lh = int(width * 0.037)

    def column(sections, x0, x1, y, head_color, text_color, decorate):
        chars = max(14, int((x1 - x0) / (width * 0.0135)))
        for sec in sections:
            if sec.heading:
                d.text((x0, y), sec.heading.upper(), font=head, fill=head_color)
                y += int(width * 0.043)
                if decorate and theme.heading == "rule":
                    d.line([x0, y - 4, x1, y - 4], fill=pal.accent, width=2)
                elif decorate and theme.heading == "bar":
                    d.rectangle([x0, y - 5, x0 + int(width * 0.06), y - 2], fill=pal.accent)
            for ln in sec.lines:
                if not ln.strip():
                    y += lh // 3
                    continue
                text = ln.strip().lstrip("-•* ").strip()
                ind = 24 if ln.startswith("  ") else 0
                for part in textwrap.wrap(text, max(10, chars - ind // 8)):
                    if y > height - 20:
                        return y
                    d.text((x0 + ind, y), part, font=small, fill=text_color)
                    y += lh
            y += int(width * 0.015)
        return y

    side_y = band_h + (int(width * 0.26) if theme.avatar else int(width * 0.06))
    column(side, sx0 + pad, sx0 + sw - int(width * 0.03), side_y, pal.sidebar_heading, pal.sidebar_text, False)
    mx0 = sw if theme.sidebar_side == "left" else 0
    mx = mx0 + pad
    my = band_h + int(width * 0.06)
    if not band_h:
        d.text((mx, my), p.name, font=name_font, fill=pal.section_title)
        if p.subtitle:
            d.text((mx, my + int(width * 0.065)), p.subtitle, font=role_font, fill=pal.section_title)
        my += int(width * 0.12)
    column(main, mx, mx0 + width - sw - int(width * 0.03), my, pal.section_title, pal.text, True)
    buf = BytesIO()
    img.save(buf, format="PNG", optimize=True)
    return buf.getvalue()
