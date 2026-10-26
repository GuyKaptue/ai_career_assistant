"""Template-aware document rendering (HTML, PDF, DOCX) from one shared theme.

Every built-in template carries ``layout`` and ``color_scheme`` metadata. This module
turns that metadata into a visual theme (accent colour, font, header and heading
style) and renders the document text in that theme. The designs are original,
inspired by common open-source CV themes (clean single-column, banded header,
timeline rule, compact startup) and use only libraries already pinned
(reportlab, python-docx, Pillow). No personal data is read or logged.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from html import escape
from io import BytesIO
from typing import Final

from app.domain.templates.themes import CvTheme

SCHEME_COLORS: Final[dict[str, str]] = {
    "blue": "#1F5FA8",
    "slate": "#334155",
    "dark": "#111827",
    "neutral": "#4B5563",
    "maroon": "#7F1D1D",
    "teal": "#0F766E",
    "coral": "#E2573F",
    "indigo": "#4338CA",
    "olive": "#4D6B2F",
    "steel": "#47627A",
    "orange": "#D9731A",
    "charcoal": "#2D2F33",
    "navy": "#0B2A55",
    "ivory": "#8A6F3B",
    "amber": "#B7791F",
    "gray": "#52525B",
    "violet": "#6D28D9",
    "cyan": "#0E7490",
}
DEFAULT_COLOR: Final = "#1F5FA8"


@dataclass(frozen=True)
class Theme:
    layout: str
    accent: str
    font: str  # reportlab base font: Helvetica | Times-Roman
    header: str  # band | left | center
    heading: str  # rule | caps | bar | plain
    compact: bool = False
    columns: int = 1  # 2 = sidebar 32% / main 68%
    palette: CvTheme | None = None  # selectable colour theme (None = template colours)
    avatar: bool = False  # draw an initials avatar slot (samples only; no photo data is stored)
    sidebar_side: str = "left"  # left | right (two-column layouts)
    sidebar_pct: int = 32
    top_band: bool = False  # full-width name band above the two columns

    @property
    def heading_color(self) -> str:
        return self.palette.section_title if self.palette else self.accent

    @property
    def rgb(self) -> tuple[int, int, int]:
        h = self.accent.lstrip("#")
        return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)

    @property
    def docx_font(self) -> str:
        return "Times New Roman" if self.font == "Times-Roman" else "Calibri"


# layout -> (font, header, heading, compact)
_LAYOUTS: Final[dict[str, tuple[str, str, str, bool]]] = {
    "classic": ("Helvetica", "left", "rule", False),
    "modern": ("Helvetica", "band", "caps", False),
    "executive": ("Times-Roman", "band", "rule", False),
    "minimal": ("Helvetica", "left", "plain", True),
    "academic": ("Times-Roman", "center", "rule", False),
    "technical": ("Helvetica", "left", "bar", True),
    "creative": ("Helvetica", "band", "bar", False),
    "timeline": ("Helvetica", "left", "bar", False),
    "functional": ("Helvetica", "left", "caps", False),
    "international": ("Helvetica", "center", "rule", False),
    "startup": ("Helvetica", "band", "caps", True),
    "corporate": ("Times-Roman", "left", "rule", False),
    "din5008": ("Helvetica", "left", "plain", False),
    "formal": ("Times-Roman", "left", "plain", False),
    "warm": ("Helvetica", "left", "plain", False),
    "brief": ("Helvetica", "left", "plain", True),
    "letterhead": ("Helvetica", "band", "plain", False),
    "centered": ("Times-Roman", "center", "plain", False),
    "modern-letter": ("Helvetica", "left", "plain", False),
    "executive-letter": ("Times-Roman", "band", "plain", False),
}


def with_palette(base: Theme, palette: CvTheme, *, columns: int = 1, avatar: bool = False) -> Theme:
    """Apply a selectable colour theme (and column count) on top of a template theme."""
    from dataclasses import replace

    font = "Times-Roman" if palette.font == "Georgia" else "Helvetica"
    return replace(base, accent=palette.accent, font=font, columns=columns, palette=palette, avatar=avatar)


def theme_from_template(
    metadata: dict,
    kind: str,
    *,
    theme_id: str = "",
    country: str = "",
    sidebar: str = "",
    locale: str = "en",
    avatar: bool | None = None,
) -> Theme:
    """Single place that turns template metadata (+ optional theme/country/sidebar overrides) into a Theme.

    A template may declare ``columns``, ``sidebar`` (left|right), ``sidebar_pct``, ``top_band``,
    ``heading`` and a default ``theme_id``. A country profile overrides the column count
    (one-column countries always get one column); explicit ``sidebar`` / ``theme_id`` win.
    """
    from dataclasses import replace

    from app.domain.documents.country_profiles import profile_for_country
    from app.domain.templates.themes import resolve_theme

    base = theme_for(metadata.get("layout"), metadata.get("color_scheme"))
    columns = int(metadata.get("columns", 1)) if kind == "cv" else 1
    show_avatar = columns == 2
    if country:
        profile = profile_for_country(country, locale)
        columns = 2 if profile.layout == "two_column" and kind == "cv" else 1
        show_avatar = columns == 2 and profile.photo != "forbidden"
    palette_id = theme_id or str(metadata.get("theme_id") or "")
    if not (palette_id or country or columns == 2):
        return base
    palette, _ = resolve_theme(palette_id or None, single_column=columns == 1)
    themed = with_palette(base, palette, columns=columns, avatar=show_avatar if avatar is None else avatar)
    if columns == 2:
        themed = replace(
            themed,
            sidebar_side=sidebar or str(metadata.get("sidebar", "left")),
            sidebar_pct=int(metadata.get("sidebar_pct", 32)),
            top_band=str(metadata.get("top_band", "")).lower() in ("true", "1"),
            heading=str(metadata.get("heading", base.heading)),
        )
    return themed


def theme_for(layout: str | None, color_scheme: str | None) -> Theme:
    font, header, heading, compact = _LAYOUTS.get(layout or "", _LAYOUTS["classic"])
    accent = SCHEME_COLORS.get(color_scheme or "", DEFAULT_COLOR)
    return Theme(layout or "classic", accent, font, header, heading, compact)


# ----------------------------------------------------------------------------- parsing

_HEADINGS: Final = {
    "professional experience",
    "experience",
    "education",
    "skills",
    "profile",
    "berufserfahrung",
    "ausbildung",
    "kenntnisse",
    "profil",
    "expérience professionnelle",
    "formation",
    "compétences",
}
_BULLET = re.compile(r"^(\s*)([-•*])\s+(.*)$")


@dataclass
class Section:
    heading: str | None
    lines: list[str]


@dataclass
class Parsed:
    name: str
    subtitle: str
    sections: list[Section]
    is_letter: bool


def _is_heading(line: str) -> bool:
    low = line.strip().rstrip(":").lower()
    return low in _HEADINGS or (line.strip().endswith(":") and len(line) < 90 and "(" in line)


def parse_content(content: str, document_type: str, title: str = "") -> Parsed:
    blocks = [b.strip("\n") for b in re.split(r"\n\s*\n", content.strip()) if b.strip()]
    if not blocks:
        return Parsed(title or "Document", "", [], document_type != "cv")
    first = blocks[0].splitlines()
    name = first[0].strip()
    rest_first = "\n".join(first[1:]).strip()
    rest = ([rest_first] if rest_first else []) + blocks[1:]
    if document_type != "cv":
        return Parsed(name, "", [Section(None, [ln for b in rest for ln in (*b.splitlines(), "")][:-1])], True)
    subtitle = ""
    if rest and "\n" not in rest[0] and not _is_heading(rest[0]):
        subtitle = rest.pop(0).strip()
    while rest and rest[0].strip() == subtitle:
        rest.pop(0)
    sections: list[Section] = []
    for block in rest:
        lines = block.splitlines()
        if _is_heading(lines[0]):
            sections.append(Section(lines[0].strip().rstrip(":"), lines[1:]))
        elif sections and sections[-1].heading is None:
            sections[-1].lines += ["", *lines]
        else:
            sections.append(Section(None, lines))
    return Parsed(name, subtitle, sections, False)


def _split_bullet(line: str) -> tuple[int, str] | None:
    m = _BULLET.match(line)
    return (1 if len(m.group(1)) >= 2 else 0, m.group(3)) if m else None


# ----------------------------------------------------------------------------- HTML


def render_styled_html(title: str, content: str, document_type: str, theme: Theme) -> str:
    if theme.columns == 2 and document_type == "cv":
        from app.infrastructure.document_export.styled_columns import html_two_column

        return html_two_column(title, content, theme)
    p = parse_content(content, document_type, title)
    body: list[str] = []
    for sec in p.sections:
        if sec.heading:
            body.append(f'<h2 class="h-{theme.heading}">{escape(sec.heading)}</h2>')
        items: list[str] = []
        for ln in sec.lines:
            b = _split_bullet(ln)
            if b:
                items.append(f'<li class="l{b[0]}">{escape(b[1])}</li>')
                continue
            if items:
                body.append("<ul>" + "".join(items) + "</ul>")
                items = []
            if ln.strip():
                body.append(f"<p>{escape(ln.strip())}</p>")
        if items:
            body.append("<ul>" + "".join(items) + "</ul>")
    sub = f'<div class="sub">{escape(p.subtitle)}</div>' if p.subtitle else ""
    gap = "6px" if theme.compact else "12px"
    family = "Georgia,'Times New Roman',serif" if theme.font == "Times-Roman" else "'Helvetica Neue',Arial,sans-serif"
    align = "center" if theme.header == "center" else "left"
    band = (
        f"background:{theme.accent};color:#fff;padding:22px 28px;"
        if theme.header == "band"
        else f"color:{theme.accent};padding:0 0 10px;"
    )
    css = (
        f"body{{margin:0;background:#f3f4f6;font-family:{family};color:#222;font-size:14px;line-height:1.45}}"
        f".page{{max-width:794px;margin:24px auto;background:#fff;padding:0 0 32px;box-shadow:0 1px 6px #0002}}"
        f".hdr{{{band}text-align:{align}}}.hdr h1{{margin:0;font-size:28px}}.sub{{opacity:.9;margin-top:4px}}"
        f".c{{padding:8px 28px}}p{{margin:{gap} 0}}ul{{margin:4px 0;padding-left:20px}}li.l1{{margin-left:18px}}"
        f"h2{{font-size:15px;margin:18px 0 6px;color:{theme.heading_color}}}"
        f".h-rule{{border-bottom:2px solid {theme.accent};padding-bottom:3px}}"
        f".h-caps{{text-transform:uppercase;letter-spacing:.08em}}"
        f".h-bar{{border-left:5px solid {theme.accent};padding-left:8px}}"
    )
    return (
        '<!DOCTYPE html>\n<html lang="en"><head><meta charset="utf-8" />'
        f'<meta name="viewport" content="width=device-width,initial-scale=1" />'
        f"<title>{escape(title or p.name)}</title><style>{css}</style></head><body>"
        f'<div class="page"><div class="hdr"><h1>{escape(p.name)}</h1>{sub}</div>'
        f'<div class="c">{"".join(body)}</div></div></body></html>\n'
    )


# ----------------------------------------------------------------------------- PDF


def render_styled_pdf(title: str, content: str, document_type: str, theme: Theme) -> bytes:
    if theme.columns == 2 and document_type == "cv":
        from app.infrastructure.document_export.styled_columns import pdf_two_column

        return pdf_two_column(title, content, theme)
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    def esc(text: str) -> str:
        return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    p = parse_content(content, document_type, title)
    accent = colors.HexColor(theme.accent)
    base = theme.font
    bold = f"{base}-Bold" if base == "Helvetica" else "Times-Bold"
    size = 9.5 if theme.compact else 10.5
    lead = size * 1.4
    align = 1 if theme.header == "center" else 0
    body = ParagraphStyle("body", fontName=base, fontSize=size, leading=lead, spaceAfter=2 if theme.compact else 4)
    sub = ParagraphStyle("bullet", parent=body, leftIndent=12, bulletIndent=2, spaceAfter=1 if theme.compact else 2)
    sub2 = ParagraphStyle("bullet2", parent=sub, leftIndent=26, bulletIndent=16)
    band_on = theme.header == "band"
    name_style = ParagraphStyle(
        "name",
        fontName=bold,
        fontSize=22,
        leading=26,
        alignment=align,
        textColor=colors.white if band_on else accent,
    )
    sub_style = ParagraphStyle(
        "sub",
        fontName=base,
        fontSize=11,
        leading=14,
        alignment=align,
        textColor=colors.white if band_on else colors.grey,
    )
    h_style = ParagraphStyle(
        "h",
        fontName=bold,
        fontSize=11.5,
        leading=14,
        textColor=colors.HexColor(theme.heading_color),
        spaceBefore=10,
        spaceAfter=3,
        leftIndent=7 if theme.heading == "bar" else 0,
    )

    story: list = []
    head = [[Paragraph(esc(p.name), name_style)]]
    if p.subtitle:
        head.append([Paragraph(esc(p.subtitle), sub_style)])
    table = Table(head, colWidths=[A4[0] - 40 * mm])
    style = [("LEFTPADDING", (0, 0), (-1, -1), 8 if band_on else 0), ("TOPPADDING", (0, 0), (-1, -1), 4)]
    if band_on:
        style += [
            ("BACKGROUND", (0, 0), (-1, -1), accent),
            ("TOPPADDING", (0, 0), (0, 0), 14),
            ("BOTTOMPADDING", (0, -1), (-1, -1), 14),
        ]
    table.setStyle(TableStyle(style))
    story += [table, Spacer(1, 8)]
    if not band_on:
        story.append(HRFlowable(width="100%", thickness=1.2, color=accent, spaceAfter=4))

    for sec in p.sections:
        if sec.heading:
            text = esc(sec.heading.upper() if theme.heading == "caps" else sec.heading)
            story.append(Paragraph(text, h_style))
            if theme.heading == "rule":
                story.append(HRFlowable(width="100%", thickness=0.8, color=accent, spaceAfter=3))
            elif theme.heading == "bar":
                story.append(HRFlowable(width=40, thickness=2.5, color=accent, spaceAfter=3, hAlign="LEFT"))
        for ln in sec.lines:
            b = _split_bullet(ln)
            if b:
                story.append(Paragraph(esc(b[1]), sub2 if b[0] else sub, bulletText="•"))
            elif ln.strip():
                story.append(Paragraph(esc(ln.strip()), body))
            else:
                story.append(Spacer(1, 4))
    buffer = BytesIO()
    SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=title or p.name,
    ).build(story)
    return buffer.getvalue()


# ----------------------------------------------------------------------------- DOCX


def _shade(cell, hex_color: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    tc_pr.append(shd)


def _bottom_border(paragraph, hex_color: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, val in (("w:val", "single"), ("w:sz", "8"), ("w:space", "1"), ("w:color", hex_color.lstrip("#"))):
        bottom.set(qn(key), val)
    borders.append(bottom)
    p_pr.append(borders)


def render_styled_docx(title: str, content: str, document_type: str, theme: Theme) -> bytes:
    if theme.columns == 2 and document_type == "cv":
        from app.infrastructure.document_export.styled_columns import docx_two_column

        return docx_two_column(title, content, theme)
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    p = parse_content(content, document_type, title)
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name = theme.docx_font
    normal.font.size = Pt(9.5 if theme.compact else 10.5)
    accent = RGBColor(*theme.rgb)
    align = WD_ALIGN_PARAGRAPH.CENTER if theme.header == "center" else WD_ALIGN_PARAGRAPH.LEFT

    if theme.header == "band":
        table = doc.add_table(rows=1, cols=1)
        cell = table.rows[0].cells[0]
        _shade(cell, theme.accent)
        para = cell.paragraphs[0]
        run = para.add_run(p.name)
        run.bold, run.font.size, run.font.color.rgb = True, Pt(22), RGBColor(255, 255, 255)
        if p.subtitle:
            sp = cell.add_paragraph()
            sr = sp.add_run(p.subtitle)
            sr.font.size, sr.font.color.rgb = Pt(11), RGBColor(255, 255, 255)
        doc.add_paragraph()
    else:
        para = doc.add_paragraph()
        para.alignment = align
        run = para.add_run(p.name)
        run.bold, run.font.size, run.font.color.rgb = True, Pt(22), accent
        if p.subtitle:
            sp = doc.add_paragraph()
            sp.alignment = align
            sr = sp.add_run(p.subtitle)
            sr.font.size = Pt(11)
        _bottom_border(sp if p.subtitle else para, theme.accent)

    for sec in p.sections:
        if sec.heading:
            hp = doc.add_paragraph()
            hp.paragraph_format.space_before = Pt(10)
            hr = hp.add_run(sec.heading.upper() if theme.heading == "caps" else sec.heading)
            hr.bold, hr.font.size, hr.font.color.rgb = (
                True,
                Pt(11.5),
                RGBColor.from_string(theme.heading_color.lstrip("#").upper()),
            )
            if theme.heading == "rule":
                _bottom_border(hp, theme.accent)
        for ln in sec.lines:
            b = _split_bullet(ln)
            if b:
                bp = doc.add_paragraph(b[1], style="List Bullet")
                if b[0]:
                    bp.paragraph_format.left_indent = Cm(1.6)
            elif ln.strip():
                doc.add_paragraph(ln.strip())
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
